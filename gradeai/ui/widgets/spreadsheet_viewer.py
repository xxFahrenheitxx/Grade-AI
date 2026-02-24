"""Excel spreadsheet (.xlsx) viewer using openpyxl."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Optional

from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QComboBox,
    QHeaderView,
    QHBoxLayout,
    QLabel,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

logger = logging.getLogger(__name__)

# openpyxl is an optional-but-expected dependency
try:
    from openpyxl import load_workbook

    _HAS_OPENPYXL = True
except ImportError:  # pragma: no cover
    _HAS_OPENPYXL = False
    logger.warning("openpyxl is not installed; .xlsx viewing is disabled.")

_DEFAULT_FONT_SIZE = 11
_MIN_FONT_SIZE = 6
_MAX_FONT_SIZE = 48


class SpreadsheetViewer(QWidget):
    """Read-only viewer for Excel ``.xlsx`` files.

    Renders the workbook using a ``QTableWidget`` and provides a sheet
    selector (``QComboBox``) when the workbook contains multiple sheets.
    """

    def __init__(self, parent=None) -> None:
        super().__init__(parent)
        self._file_path: Optional[Path] = None
        self._workbook = None
        self._font_size: int = _DEFAULT_FONT_SIZE

        # Layout
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)

        # Sheet selector toolbar
        self._toolbar = QWidget(self)
        toolbar_layout = QHBoxLayout(self._toolbar)
        toolbar_layout.setContentsMargins(4, 4, 4, 4)
        toolbar_layout.setSpacing(6)

        self._sheet_label = QLabel("Sheet:")
        toolbar_layout.addWidget(self._sheet_label)

        self._sheet_combo = QComboBox()
        self._sheet_combo.setMinimumWidth(150)
        self._sheet_combo.currentIndexChanged.connect(self._on_sheet_changed)
        toolbar_layout.addWidget(self._sheet_combo)
        toolbar_layout.addStretch()

        layout.addWidget(self._toolbar)

        # Table
        self._table = QTableWidget(self)
        self._table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self._table.horizontalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.ResizeToContents
        )
        self._table.verticalHeader().setSectionResizeMode(
            QHeaderView.ResizeMode.ResizeToContents
        )
        layout.addWidget(self._table)

        # Apply initial font
        font = self._table.font()
        font.setPointSize(self._font_size)
        self._table.setFont(font)

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------

    def load_file(self, path: Path) -> None:
        """Open an ``.xlsx`` file and populate the table."""
        self._file_path = path
        self._table.clear()
        self._sheet_combo.blockSignals(True)
        self._sheet_combo.clear()

        if not _HAS_OPENPYXL:
            self._table.setRowCount(1)
            self._table.setColumnCount(1)
            item = QTableWidgetItem("openpyxl is required to view .xlsx files.")
            self._table.setItem(0, 0, item)
            self._sheet_combo.blockSignals(False)
            return

        try:
            self._workbook = load_workbook(str(path), read_only=True, data_only=True)
        except Exception:
            logger.exception("Failed to open xlsx: %s", path)
            self._table.setRowCount(1)
            self._table.setColumnCount(1)
            item = QTableWidgetItem("Error: could not open spreadsheet.")
            self._table.setItem(0, 0, item)
            self._sheet_combo.blockSignals(False)
            return

        sheet_names = self._workbook.sheetnames
        self._sheet_combo.addItems(sheet_names)
        self._sheet_combo.blockSignals(False)

        # Show/hide the sheet selector toolbar based on sheet count
        self._toolbar.setVisible(len(sheet_names) > 1)

        # Load the first sheet
        if sheet_names:
            self._load_sheet(sheet_names[0])

    @property
    def file_path(self) -> Optional[Path]:
        return self._file_path

    # ------------------------------------------------------------------
    # Sheet loading
    # ------------------------------------------------------------------

    def _on_sheet_changed(self, index: int) -> None:
        if self._workbook is None or index < 0:
            return
        sheet_name = self._sheet_combo.itemText(index)
        self._load_sheet(sheet_name)

    def _load_sheet(self, sheet_name: str) -> None:
        """Populate the table widget with data from the named sheet."""
        if self._workbook is None:
            return

        ws = self._workbook[sheet_name]

        # Determine dimensions
        rows = list(ws.iter_rows())
        if not rows:
            self._table.setRowCount(0)
            self._table.setColumnCount(0)
            return

        row_count = len(rows)
        col_count = max(len(row) for row in rows) if rows else 0

        self._table.setRowCount(row_count)
        self._table.setColumnCount(col_count)

        for r_idx, row in enumerate(rows):
            for c_idx, cell in enumerate(row):
                value = cell.value
                text = str(value) if value is not None else ""
                item = QTableWidgetItem(text)
                item.setFlags(item.flags() & ~Qt.ItemFlag.ItemIsEditable)
                self._table.setItem(r_idx, c_idx, item)

    # ------------------------------------------------------------------
    # CTRL+scroll zoom
    # ------------------------------------------------------------------

    def wheelEvent(self, event) -> None:  # noqa: N802
        if event.modifiers() & Qt.KeyboardModifier.ControlModifier:
            angle_delta = event.angleDelta().y()
            if angle_delta > 0:
                self._font_size = min(_MAX_FONT_SIZE, self._font_size + 1)
            elif angle_delta < 0:
                self._font_size = max(_MIN_FONT_SIZE, self._font_size - 1)
            else:
                return
            font = self._table.font()
            font.setPointSize(self._font_size)
            self._table.setFont(font)
            # Also update header fonts
            self._table.horizontalHeader().setFont(font)
            self._table.verticalHeader().setFont(font)
            event.accept()
        else:
            super().wheelEvent(event)
