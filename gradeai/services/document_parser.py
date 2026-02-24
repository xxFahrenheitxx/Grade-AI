"""Document content extraction for AI grading context."""

from __future__ import annotations

import base64
import logging
from pathlib import Path
from typing import Any, Optional

from gradeai.utils.file_utils import get_file_category

logger = logging.getLogger(__name__)


def extract_content_blocks(file_path: Path) -> list[dict[str, Any]]:
    """Extract content from a file as multi-modal content blocks for the AI API.

    Returns a list of content blocks (text and/or image) suitable for
    inclusion in an Anthropic API message.
    """
    category = get_file_category(file_path)

    if category == "pdf":
        return _extract_pdf(file_path)
    elif category == "image":
        return _extract_image(file_path)
    elif category == "word":
        return _extract_docx(file_path)
    elif category == "excel":
        return _extract_xlsx(file_path)
    elif category in ("code", "text", "markdown"):
        return _extract_text(file_path)
    else:
        logger.warning("Unsupported file type for content extraction: %s", file_path.suffix)
        return [{"type": "text", "text": f"[Unsupported file type: {file_path.name}]"}]


def _extract_pdf(file_path: Path) -> list[dict[str, Any]]:
    """Extract text and page images from a PDF file."""
    import fitz  # PyMuPDF

    blocks: list[dict[str, Any]] = []
    with fitz.open(str(file_path)) as doc:
        for page_num in range(len(doc)):
            page = doc.load_page(page_num)

            # Render page as image for visual analysis (handwriting, diagrams, layout)
            pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0))
            img_bytes = pix.tobytes("png")
            b64 = base64.b64encode(img_bytes).decode("ascii")
            blocks.append({
                "type": "image",
                "source": {
                    "type": "base64",
                    "media_type": "image/png",
                    "data": b64,
                },
            })

            # Also extract machine-readable text
            text = page.get_text()
            if text.strip():
                blocks.append({
                    "type": "text",
                    "text": f"--- Page {page_num + 1} text ---\n{text.strip()}",
                })

    return blocks


def _extract_image(file_path: Path) -> list[dict[str, Any]]:
    """Extract an image file as a base64 content block."""
    suffix = file_path.suffix.lower()
    media_types = {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".gif": "image/gif",
        ".bmp": "image/bmp",
        ".webp": "image/webp",
        ".tiff": "image/tiff",
    }
    media_type = media_types.get(suffix, "image/png")

    img_bytes = file_path.read_bytes()
    b64 = base64.b64encode(img_bytes).decode("ascii")

    return [{
        "type": "image",
        "source": {
            "type": "base64",
            "media_type": media_type,
            "data": b64,
        },
    }]


def _extract_docx(file_path: Path) -> list[dict[str, Any]]:
    """Extract text content from a Word document."""
    from docx import Document

    doc = Document(str(file_path))
    paragraphs = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text)

    # Also extract table content
    for table in doc.tables:
        table_text = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            table_text.append(" | ".join(cells))
        if table_text:
            paragraphs.append("\n[Table]\n" + "\n".join(table_text))

    text = "\n\n".join(paragraphs)
    return [{"type": "text", "text": text}]


def _extract_xlsx(file_path: Path) -> list[dict[str, Any]]:
    """Extract content from an Excel spreadsheet."""
    from openpyxl import load_workbook

    wb = load_workbook(str(file_path), data_only=True)
    blocks: list[dict[str, Any]] = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows_text = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(cell) if cell is not None else "" for cell in row]
            if any(c.strip() for c in cells):
                rows_text.append(" | ".join(cells))

        if rows_text:
            sheet_text = f"--- Sheet: {sheet_name} ---\n" + "\n".join(rows_text)
            blocks.append({"type": "text", "text": sheet_text})

    wb.close()
    return blocks if blocks else [{"type": "text", "text": "[Empty spreadsheet]"}]


def _extract_text(file_path: Path) -> list[dict[str, Any]]:
    """Extract plain text content from a text-based file."""
    try:
        text = file_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        text = file_path.read_text(encoding="latin-1")

    return [{"type": "text", "text": text}]


def build_grading_context(
    exam_path: Path,
    solution_paths: list[Path],
    rule_paths: list[Path],
    exam_paper_path: Optional[Path] = None,
) -> list[dict[str, Any]]:
    """Build the complete multi-modal content for an AI grading request.

    Combines exam paper (if any), exam content, solution documents,
    and grading rules into a single list of content blocks.
    The exam paper is placed first for highest AI context priority.
    """
    blocks: list[dict[str, Any]] = []

    # Official exam paper (highest priority -- placed first)
    if exam_paper_path and exam_paper_path.exists():
        blocks.append({"type": "text", "text": "## Official Exam Paper (Questions)\n"})
        blocks.extend(extract_content_blocks(exam_paper_path))

    # Student exam content
    blocks.append({"type": "text", "text": "\n\n## Student Exam\n"})
    blocks.extend(extract_content_blocks(exam_path))

    # Solution documents
    if solution_paths:
        blocks.append({"type": "text", "text": "\n\n## Solution Documents\n"})
        for sol_path in solution_paths:
            blocks.append({"type": "text", "text": f"\n### {sol_path.name}\n"})
            blocks.extend(extract_content_blocks(sol_path))

    # Grading rules
    if rule_paths:
        blocks.append({"type": "text", "text": "\n\n## Grading Rules\n"})
        for rule_path in rule_paths:
            blocks.append({"type": "text", "text": f"\n### {rule_path.name}\n"})
            blocks.extend(extract_content_blocks(rule_path))

    return blocks
